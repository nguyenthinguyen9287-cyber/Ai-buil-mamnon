import streamlit as st
import os
from dotenv import load_dotenv
from docx import Document
from io import BytesIO
import google.generativeai as genai

load_dotenv()
API_KEY = os.getenv("GOOGLE_API_KEY")

if not API_KEY:
    st.error("Chua tim thay API Key!")
    st.stop()

genai.configure(api_key=API_KEY)
model = genai.GenerativeModel("gemini-3.8-flash")

def goi_ai(prompt):
    try:
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        return f"Loi: {str(e)}"

st.set_page_config(page_title="AI-Build Mam Non", layout="wide")
st.title("AI-BUILD MẦM NON")
st.subheader("Soạn giáo án & tạo trò chơi học tập")
st.divider()

tab1, tab2 = st.tabs(["📋 Soạn Giáo Án", "🎮 Tạo Trò Chơi"])

with tab1:
    st.header("Soạn Giáo Án")
    col1, col2 = st.columns(2)
    with col1:
        chu_de = st.text_input("🌸 Tên chủ đề/bài học:")
        tuoi = st.selectbox("👶 Độ tuổi", ["18-24 tháng", "25-36 tháng", "3-4 tuổi", "4-5 tuổi", "5-6 tuổi"])
        linh_vuc = st.multiselect("🧠 Lĩnh vực phát triển",
            ["Thể chất", "Ngôn ngữ", "Nhận thức", "Tình cảm - Xã hội", "Thẩm mỹ"],
            default=["Nhận thức", "Ngôn ngữ"])
    with col2:
        hinh_thuc = st.selectbox("📝 Hình thức tổ chức",
            ["Hoạt động theo chủ đề", "Hoạt động trải nghiệm", "Hoạt động ngoài trời", "Hoạt động góc"])
        thoi_gian = st.slider("⏱️ Thời lượng (phút)", 15, 45, 30)
        ghi_chu = st.text_area("💡 Yêu cầu/Ghi chú thêm:", height=100)

    if st.button("✨ Tạo Giáo Án", type="primary", use_container_width=True) and chu_de:
        with st.spinner("🤖 AI đang soạn giáo án..."):
            prompt = f"""Soạn giáo án chi tiết cho mầm non theo Chương trình Giáo dục Mầm non Việt Nam.

Chủ đề/bài học: {chu_de}
Độ tuổi: {tuoi}
Lĩnh vực phát triển: {', '.join(linh_vuc)}
Hình thức tổ chức: {hinh_thuc}
Thời lượng: {thoi_gian} phút
Yêu cầu thêm: {ghi_chu}

Trình bày rõ cấu trúc sau:

I. MỤC TIÊU
- Mục tiêu chung
- Mục tiêu cụ thể (Kiến thức - Kỹ năng - Thái độ)

II. CHUẨN BỊ
- Cơ sở vật chất
- Chuẩn bị của giáo viên
- Chuẩn bị của trẻ

III. TIẾN TRÌNH HOẠT ĐỘNG
1. Hoạt động mở đầu
2. Hoạt động phát triển
3. Hoạt động kết thúc

IV. GỢI Ý HÌNH ẢNH MINH HỌA

V. LƯU Ý CHO GIÁO VIÊN

Viết ngắn gọn, dễ hiểu, phù hợp với độ tuổi và điều kiện thực tế tại trường mầm non Việt Nam."""
            kq = goi_ai(prompt)
        
        if "Loi:" not in kq:
            st.success("✅ Soạn giáo án xong!")
            st.markdown(kq)
            doc = Document()
            doc.add_heading(f"GIÁO ÁN: {chu_de} – {tuoi}", 0)
            doc.add_paragraph(kq)
            tep = BytesIO()
            doc.save(tep)
            tep.seek(0)
            st.download_button("📥 Tải file Word", data=tep, file_name=f"GiaoAn_{chu_de}.docx")
        else:
            st.error(kq)
    
    st.info("💡 Giáo viên vui lòng xem xét & điều chỉnh nội dung cho phù hợp với lớp học nhé!")

with tab2:
    st.header("Tạo Trò Chơi Học Tập")
    col1, col2 = st.columns(2)
    with col1:
        ten_tro_choi = st.text_input("🎮 Tên trò chơi:")
        tuoi2 = st.selectbox("👶 Độ tuổi", ["18-24 tháng", "25-36 tháng", "3-4 tuổi", "4-5 tuổi", "5-6 tuổi"], key="t2")
        muc_tieu_tro = st.multiselect("🎯 Mục tiêu trò chơi",
            ["Phát triển ngôn ngữ", "Nhận biết màu sắc", "Nhận biết số lượng", "Phát triển vận động", "Rèn luyện kỹ năng xã hội"],
            default=["Nhận biết màu sắc"])
    with col2:
        so_nguoi = st.selectbox("👥 Số lượng tham gia", ["Cá nhân", "Nhóm nhỏ (3-5 bé)", "Toàn lớp"])
        ghi_chu_tro = st.text_area("💡 Ý tưởng/Ghi chú thêm:", height=100)

    if st.button("✨ Tạo Trò Chơi", type="primary", use_container_width=True) and ten_tro_choi:
        with st.spinner("🤖 AI đang thiết kế trò chơi..."):
            prompt = f"""Thiết kế trò chơi học tập cho trẻ mầm non.

Tên trò chơi: {ten_tro_choi}
Độ tuổi: {tuoi2}
Mục tiêu: {', '.join(muc_tieu_tro)}
Số lượng tham gia: {so_nguoi}
Yêu cầu thêm: {ghi_chu_tro}

Trình bày rõ:

I. THÔNG TIN CHUNG
- Tên trò chơi
- Đối tượng
- Thời lượng dự kiến

II. MỤC TIÊU
- Mục tiêu giáo dục
- Giá trị giáo dục

III. CHUẨN BỊ
- Vật liệu, đồ dùng
- Không gian tổ chức

IV. CÁCH TIẾN HÀNH
1. Giới thiệu trò chơi
2. Quy tắc trò chơi
3. Các bước thực hiện
4. Kết thúc trò chơi

V. GỢI Ý BIẾN THỂ & LƯU Ý

Ngôn ngữ đơn giản, dễ thực hiện, phù hợp với điều kiện lớp học mầm non Việt Nam."""
            kq = goi_ai(prompt)
        
        if "Loi:" not in kq:
            st.success("✅ Thiết kế trò chơi xong!")
            st.markdown(kq)
            doc = Document()
            doc.add_heading(f"TRÒ CHƠI: {ten_tro_choi} – {tuoi2}", 0)
            doc.add_paragraph(kq)
            tep = BytesIO()
            doc.save(tep)
            tep.seek(0)
            st.download_button("📥 Tải file Word", data=tep, file_name=f"TroChoi_{ten_tro_choi}.docx")
        else:
            st.error(kq)
    
    st.info("💡 Hãy điều chỉnh trò chơi cho phù hợp với điều kiện và đặc điểm trẻ ở lớp nhé!")

st.divider()
st.caption("© 2026 AI-Build Mầm Non")